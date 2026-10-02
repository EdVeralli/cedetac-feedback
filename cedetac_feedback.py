# -*- coding: utf-8 -*-
'''
CEDETAC - Feedback de usuarios (encuesta CXF de Boti)

Flujo:
  1. Lee el período de config_fechas.txt (o de --desde/--hasta)
  2. Pide el login AWS vía aws-azure-login y verifica credenciales
  3. Ejecuta sql/cedetac_feedback.sql en Athena
  4. Limpia la bajada (códigos CXF, columnas, espacios)
  5. Genera el Excel de indicadores a partir de plantilla/ (fórmulas del Excel testigo)

Uso (desde la carpeta del proyecto):
  python cedetac_feedback.py
  python cedetac_feedback.py --desde 2026-07-01 --hasta 2026-09-30
  python cedetac_feedback.py --csv bajada02.csv      (reprocesa un CSV, sin Athena)

Workgroup: Production-caba-piba-athena-boti-group
Rol: PIBADataScientist
'''
import argparse
import logging
import sys
import time
from calendar import monthrange
from copy import copy
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd

# ==================== CONFIGURACION ====================
BASE = Path(__file__).resolve().parent

CONFIG = {
    'region': 'us-east-1',
    'workgroup': 'Production-caba-piba-athena-boti-group',
    'database': 'caba-piba-consume-zone-db',
    'perfil_aws': 'default',
    'config_fechas': BASE / 'config_fechas.txt',
    'sql': BASE / 'sql' / 'cedetac_feedback.sql',
    'plantilla': BASE / 'plantilla' / 'CEDETAC_plantilla.xlsx',
    'salida': BASE / 'output',
}

CMD_LOGIN = 'aws-azure-login --profile {perfil} --mode=gui'
MAX_REINTENTOS_QUERY = 3
AVISO_SEGUNDOS = 30      # cada cuánto informar el estado de la query
POLL_SEGUNDOS = 3        # cada cuánto consultar a Athena

HOJA_IND = 'Indicadores generales'
HOJA_DAT = 'Feedback CEDETAC'

# Limpieza
COLS_ELIMINAR = ['session_id', 'usuario', 'ultimotemanps', 'ultimotema', 'link_conversacion']
COLS_LIMPIAR = ['rule_name', 'esfuerzo', 'satisfaccion']
PATRON_CODIGOS = r'^\s*CXF\w+\s+|\s+(?:CATs|Integraciones|Estáticos|Pushes)\s*$'
COLS_FINALES = ['fecha', 'efectividad', 'esfuerzo', 'satisfaccion', 'sugerencia']

log = logging.getLogger('cedetac')


# ==================== FECHAS ====================

def leer_config_fechas(archivo):
    '''
    Lee config_fechas.txt. Prioridad:
      MODO 2: FECHA_INICIO + FECHA_FIN (rango)
      MODO 3: ULTIMOS_MESES=N (últimos N meses cerrados)
      MODO 1: MES + AÑO (mes completo)
    Retorna (fecha_inicio, fecha_fin) como date.
    '''
    if not Path(archivo).exists():
        raise FileNotFoundError('No se encuentra {}'.format(archivo))

    valores = {}
    with open(archivo, 'r', encoding='utf-8') as f:
        for linea in f:
            linea = linea.strip()
            if not linea or linea.startswith('#') or '=' not in linea:
                continue
            clave, valor = linea.split('=', 1)
            valores[clave.strip().upper().replace('Ñ', 'N')] = valor.strip()

    if 'FECHA_INICIO' in valores and 'FECHA_FIN' in valores:
        ini = datetime.strptime(valores['FECHA_INICIO'], '%Y-%m-%d').date()
        fin = datetime.strptime(valores['FECHA_FIN'], '%Y-%m-%d').date()
        log.info('Modo: RANGO PERSONALIZADO')
    elif 'ULTIMOS_MESES' in valores:
        n = int(valores['ULTIMOS_MESES'])
        ini, fin = ultimos_meses_cerrados(n, date.today())
        log.info('Modo: ULTIMOS %s MESES CERRADOS', n)
    elif 'MES' in valores and 'ANO' in valores:
        mes, anio = int(valores['MES']), int(valores['ANO'])
        if not 1 <= mes <= 12:
            raise ValueError('Mes inválido: {}'.format(mes))
        ini = date(anio, mes, 1)
        fin = date(anio, mes, monthrange(anio, mes)[1])
        log.info('Modo: MES COMPLETO')
    else:
        raise ValueError('config_fechas.txt no tiene una configuración válida')

    if ini > fin:
        raise ValueError('FECHA_INICIO no puede ser posterior a FECHA_FIN')
    return ini, fin


def ultimos_meses_cerrados(n, hoy):
    '''Del día 1 de hace n meses al último día del mes anterior a hoy.'''
    anio, mes = hoy.year, hoy.month - 1
    if mes == 0:
        anio, mes = anio - 1, 12
    fin = date(anio, mes, monthrange(anio, mes)[1])
    total = anio * 12 + (mes - 1) - (n - 1)
    ini = date(total // 12, total % 12 + 1, 1)
    return ini, fin


# ==================== LOGIN AWS ====================

def sesion_aws(perfil):
    '''Sesión boto3 nueva (relee ~/.aws/credentials).'''
    import boto3
    return boto3.Session(profile_name=perfil, region_name=CONFIG['region'])


def credenciales_validas(perfil):
    try:
        ident = sesion_aws(perfil).client('sts').get_caller_identity()
        log.info('Credenciales OK: %s', ident.get('Arn', ''))
        return True
    except Exception as e:
        log.warning('Credenciales no válidas: %s', str(e).splitlines()[0])
        return False


def login_aws(perfil):
    '''
    Pide al usuario que ejecute aws-azure-login en OTRA terminal y
    espera ENTER. Recién con credenciales válidas devuelve la sesión.
    '''
    comando = CMD_LOGIN.format(perfil=perfil)
    while True:
        print('')
        print('=' * 70)
        print('LOGIN AWS (rol PIBADataScientist)')
        print('=' * 70)
        print('1. Abrí OTRA terminal y ejecutá:')
        print('')
        print('     ' + comando)
        print('')
        print('2. Completá el login en el navegador y elegí el rol PIBADataScientist')
        print('3. Cuando veas "Assuming role..." volvé acá')
        print('=' * 70)
        resp = input('ENTER para continuar (q para salir): ').strip().lower()
        if resp == 'q':
            log.info('Cancelado por el usuario')
            sys.exit(1)
        time.sleep(2)
        if credenciales_validas(perfil):
            return sesion_aws(perfil)
        print('[!] El login todavía no es válido. Repetí el comando y probá de nuevo.')


# ==================== EXTRACCION ====================

def armar_query(archivo_sql, ini, fin):
    sql = Path(archivo_sql).read_text(encoding='utf-8')
    if '{fecha_inicio}' not in sql or '{fecha_fin}' not in sql:
        raise ValueError('La query no tiene los parámetros {fecha_inicio}/{fecha_fin}')
    return (sql.replace('{fecha_inicio}', ini.isoformat())
               .replace('{fecha_fin}', fin.isoformat()))


def _es_token_vencido(e):
    msg = str(e)
    return 'ExpiredToken' in msg or 'expired' in msg.lower()


def ejecutar_query(sql, perfil, session):
    '''
    Lanza la query en Athena y hace polling: cada AVISO_SEGUNDOS informa
    estado, tiempo transcurrido y datos escaneados. Si el token vence
    mientras espera, pide login y sigue esperando la MISMA query.
    Ctrl+C cancela la query en Athena.
    '''
    import awswrangler as wr
    for intento in range(1, MAX_REINTENTOS_QUERY + 1):
        try:
            athena = session.client('athena')
            qid = athena.start_query_execution(
                QueryString=sql,
                QueryExecutionContext={'Database': CONFIG['database']},
                WorkGroup=CONFIG['workgroup'],
            )['QueryExecutionId']
            log.info('Query lanzada en Athena (intento %s) | id: %s', intento, qid)
            break
        except Exception as e:
            if _es_token_vencido(e):
                log.warning('Token AWS vencido')
                session = login_aws(perfil)
                continue
            raise
    else:
        raise RuntimeError('No se pudo lanzar la query tras {} intentos'.format(MAX_REINTENTOS_QUERY))

    t0 = time.time()
    ultimo_aviso = 0
    try:
        while True:
            try:
                q = session.client('athena').get_query_execution(QueryExecutionId=qid)['QueryExecution']
            except Exception as e:
                if _es_token_vencido(e):
                    log.warning('Token AWS vencido durante la espera (la query sigue en Athena)')
                    session = login_aws(perfil)
                    continue
                raise
            estado = q['Status']['State']
            seg = time.time() - t0
            if estado in ('SUCCEEDED', 'FAILED', 'CANCELLED'):
                break
            if seg - ultimo_aviso >= AVISO_SEGUNDOS:
                mb = q.get('Statistics', {}).get('DataScannedInBytes', 0) / 1024 ** 2
                log.info('  ... %s | %d min %02d s | %.0f MB escaneados',
                         estado, seg // 60, seg % 60, mb)
                ultimo_aviso = seg
            time.sleep(POLL_SEGUNDOS)
    except KeyboardInterrupt:
        log.warning('Interrumpido: cancelando la query en Athena...')
        session.client('athena').stop_query_execution(QueryExecutionId=qid)
        sys.exit(1)

    if estado != 'SUCCEEDED':
        motivo = q['Status'].get('StateChangeReason', '')
        raise RuntimeError('La query terminó en {}: {}'.format(estado, motivo))

    mb = q['Statistics'].get('DataScannedInBytes', 0) / 1024 ** 2
    log.info('Query OK en %d min %02d s | %.0f MB escaneados', seg // 60, seg % 60, mb)

    ruta = q['ResultConfiguration']['OutputLocation']
    df = wr.s3.read_csv(ruta, boto3_session=session, dtype=str,
                        keep_default_na=False, na_values=[''])
    log.info('Resultado descargado: %s filas', len(df))
    return df


# ==================== LIMPIEZA ====================

def limpiar(df, ini, fin):
    '''Replica la limpieza validada contra el Excel testigo. No elimina filas por contenido.'''
    df = df.copy()
    df.columns = [c.lower() for c in df.columns]
    df['fecha'] = pd.to_datetime(df['fecha'])

    fuera = ~df['fecha'].between(pd.Timestamp(ini), pd.Timestamp(fin))
    if fuera.any():
        log.info('Se excluyen %s filas con fecha fuera de %s / %s', int(fuera.sum()), ini, fin)
        df = df[~fuera]

    df = df.drop(columns=[c for c in COLS_ELIMINAR if c in df.columns])
    for c in COLS_LIMPIAR:
        df[c] = df[c].astype('string').str.replace(PATRON_CODIGOS, '', regex=True).str.strip()

    sug = df['feedbacksugerencia'].astype('string')
    df['sugerencia'] = (sug.str.replace(r'[ \t]+', ' ', regex=True)
                           .str.replace(r' *\n *', '\n', regex=True)
                           .str.strip())
    df = df.rename(columns={'rule_name': 'efectividad'})
    df = df.sort_values('fecha', kind='stable').reset_index(drop=True)
    return df[COLS_FINALES]


# ==================== EXCEL ====================

def generar_excel(df, plantilla, salida, ini, fin):
    wb = openpyxl.load_workbook(plantilla)
    ws = wb[HOJA_DAT]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row)

    estilo_fecha = 'mm-dd-yy'
    for i, fila in enumerate(df.itertuples(index=False), start=2):
        for j, v in enumerate(fila, start=1):
            if pd.isna(v):
                v = None
            elif j == 1:
                v = v.to_pydatetime()
            c = ws.cell(row=i, column=j, value=v)
            if j == 1:
                c.number_format = estilo_fecha
    ws.auto_filter.ref = 'A1:E{}'.format(len(df) + 1)

    wi = wb[HOJA_IND]
    D = "'{}'!".format(HOJA_DAT)
    wi['A1'] = 'Período {:%d/%m/%Y} al {:%d/%m/%Y}'.format(ini, fin)
    wi['A3'] = 'TOTAL DEL PERÍODO'
    # Fórmulas idénticas al Excel testigo
    wi['A5'] = '=COUNTA({0}A:A)'.format(D)
    wi['B5'] = '=(COUNTIF({0}B:B,"Sí"))/((COUNTA({0}B:B)))'.format(D)
    wi['C5'] = ('=(COUNTIF({0}C:C,"Muy difícil")*5+COUNTIF({0}C:C,"Difícil")*4'
                '+COUNTIF({0}C:C,"Más o menos")*3+COUNTIF({0}C:C,"Fácil")*2'
                '+COUNTIF({0}C:C,"Muy fácil")*1)/COUNTA({0}C:C)').format(D)
    wi['D5'] = ('=(COUNTIF({0}D:D,"Muy conforme")+COUNTIF({0}D:D,"Conforme "))'
                '/COUNTA({0}D:D)').format(D)
    wi['A7'] = 'Los valores corresponden a usuarios que respondieron el feedback'

    try:
        wb.save(salida)
    except PermissionError:
        log.error('No se pudo guardar %s: cerralo si está abierto en Excel y volvé a correr', salida)
        raise


def kpis_control(df):
    '''Mismo cálculo que las fórmulas del Excel (el +1 es el encabezado que cuenta CONTARA).'''
    n = len(df) + 1
    pesos = {'Muy difícil': 5, 'Difícil': 4, 'Más o menos': 3, 'Fácil': 2, 'Muy fácil': 1}
    return {
        'Total_Usuarios': n,
        'TasaEfectividad': (df['efectividad'] == 'Sí').sum() / n,
        'Esfuerzo': df['esfuerzo'].map(pesos).sum() / n,
        'Satisfaccion': ((df['satisfaccion'] == 'Muy conforme').sum()
                         + (df['satisfaccion'] == 'Conforme ').sum()) / n,
    }


# ==================== MAIN ====================

def configurar_log(carpeta):
    carpeta.mkdir(parents=True, exist_ok=True)
    archivo = carpeta / 'cedetac_{:%Y%m%d_%H%M%S}.log'.format(datetime.now())
    logging.basicConfig(
        level=logging.INFO,
        format='%(asctime)s [%(levelname)s] %(message)s',
        datefmt='%H:%M:%S',
        handlers=[logging.StreamHandler(sys.stdout),
                  logging.FileHandler(archivo, encoding='utf-8')],
    )
    return archivo


def parse_args():
    p = argparse.ArgumentParser(description='CEDETAC - bajada de feedback de Athena + Excel de indicadores')
    p.add_argument('--desde', help='Fecha inicio YYYY-MM-DD (pisa config_fechas.txt)')
    p.add_argument('--hasta', help='Fecha fin YYYY-MM-DD (pisa config_fechas.txt)')
    p.add_argument('--csv', help='Procesar un CSV ya bajado en lugar de consultar Athena')
    p.add_argument('--config', default=str(CONFIG['config_fechas']), help='Archivo de fechas')
    p.add_argument('--perfil', default=CONFIG['perfil_aws'], help='Perfil AWS (default: default)')
    p.add_argument('--salida', default=str(CONFIG['salida']), help='Carpeta de salida')
    return p.parse_args()


def main():
    args = parse_args()
    salida = Path(args.salida)
    archivo_log = configurar_log(salida / 'logs')
    log.info('CEDETAC - Feedback | log: %s', archivo_log)

    # 1. Período
    if args.desde or args.hasta:
        if not (args.desde and args.hasta):
            sys.exit('Usar --desde y --hasta juntos')
        ini = datetime.strptime(args.desde, '%Y-%m-%d').date()
        fin = datetime.strptime(args.hasta, '%Y-%m-%d').date()
        log.info('Modo: RANGO POR PARAMETRO')
    else:
        ini, fin = leer_config_fechas(args.config)
    log.info('Período: %s al %s', ini, fin)
    sufijo = '{:%Y-%m-%d}_al_{:%Y-%m-%d}'.format(ini, fin)

    # 2-3. Bajada
    if args.csv:
        log.info('Leyendo CSV existente: %s', args.csv)
        crudo = pd.read_csv(args.csv, dtype=str, encoding='utf-8-sig')
    else:
        session = login_aws(args.perfil)
        sql = armar_query(CONFIG['sql'], ini, fin)
        crudo = ejecutar_query(sql, args.perfil, session)
        archivo_crudo = salida / 'cedetac_bajada_{}.csv'.format(sufijo)
        crudo.to_csv(archivo_crudo, index=False, encoding='utf-8-sig')
        log.info('Bajada cruda: %s', archivo_crudo)

    if crudo.empty:
        log.warning('La bajada no trajo filas para el período. Fin.')
        return

    # 4. Limpieza
    limpio = limpiar(crudo, ini, fin)
    archivo_limpio = salida / 'cedetac_limpio_{}.csv'.format(sufijo)
    limpio.rename(columns={'efectividad': 'efectividad', 'esfuerzo': 'Esfuerzo',
                           'satisfaccion': 'Satisfaccion',
                           'sugerencia': 'Sugerencia del usuario'}) \
          .assign(fecha=limpio['fecha'].dt.date) \
          .to_csv(archivo_limpio, index=False, encoding='utf-8-sig')
    log.info('CSV limpio: %s (%s filas)', archivo_limpio, len(limpio))

    # 5. Excel
    archivo_xlsx = salida / 'CEDETAC - Feedback {:%d-%m-%Y} al {:%d-%m-%Y}.xlsx'.format(ini, fin)
    generar_excel(limpio, CONFIG['plantilla'], archivo_xlsx, ini, fin)
    log.info('Excel: %s', archivo_xlsx)

    k = kpis_control(limpio)
    log.info('-' * 50)
    log.info('Total_Usuarios : %s', k['Total_Usuarios'])
    log.info('TasaEfectividad: %.2f%%', k['TasaEfectividad'] * 100)
    log.info('Esfuerzo       : %.2f', k['Esfuerzo'])
    log.info('Satisfacción %% : %.2f%%', k['Satisfaccion'] * 100)
    log.info('-' * 50)
    log.info('PROCESO COMPLETADO')


if __name__ == '__main__':
    main()
